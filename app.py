import streamlit as st
import yfinance as yf
import pandas as pd
import io
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
import logging

logging.basicConfig(level=logging.INFO)

st.set_page_config(page_title="Valuation IDC Tool", page_icon="📈", layout="centered")

st.title("📈 Company Valuation Tool (IDC)")
st.markdown("""
Esta aplicación permite descargar los datos financieros de una empresa desde Yahoo Finance
y volcarlos automáticamente en la plantilla Excel de valoración IDC.

**Instrucciones:**
1. Introduce el Ticker de la empresa (ej: AAPL, MSFT, GOOG).
2. Haz clic en "Generar Excel".
3. Descarga el archivo generado. **Nota:** Dado que los nombres de las partidas en Yahoo Finance
   difieren de TIKR, es posible que necesites ajustar manualmente las celdas en azul/naranja en las pestañas de cálculo.
""")

ticker_symbol = st.text_input("Ticker de la Empresa (ej: AAPL)", value="AAPL")

if st.button("Generar Excel", type="primary"):
    ticker_symbol = ticker_symbol.strip().upper()
    if ticker_symbol:
        try:
            with st.spinner(f'Descargando datos para {ticker_symbol} desde yfinance...'):
                ticker = yf.Ticker(ticker_symbol)

                # Fetching the data
                is_df = ticker.income_stmt
                bs_df = ticker.balance_sheet
                cf_df = ticker.cashflow

                if is_df.empty and bs_df.empty and cf_df.empty:
                    st.error(f"No se han encontrado datos para el ticker: {ticker_symbol}. Verifica que es correcto.")
                    st.stop()

                # Sorting columns chronologically (oldest to newest)
                if not is_df.empty: is_df = is_df[is_df.columns[::-1]]
                if not bs_df.empty: bs_df = bs_df[bs_df.columns[::-1]]
                if not cf_df.empty: cf_df = cf_df[cf_df.columns[::-1]]

            with st.spinner('Procesando la plantilla Excel...'):
                template_path = "Módulo 7_ Plantilla Valoración IDC v2024.3_ TIKR(8).xlsx"

                try:
                    wb = openpyxl.load_workbook(template_path)
                except Exception as e:
                    st.error(f"Error al cargar la plantilla: {e}")
                    st.stop()

                def write_df_to_sheet(wb, sheet_name, df):
                    if sheet_name in wb.sheetnames:
                        ws = wb[sheet_name]
                        # Ensure we don't have overlapping old data by clearing the sheet first?
                        # Or just overwrite. Overwriting is usually fine since the template is mostly empty.

                        df = df.reset_index()

                        # Date formatting for the columns
                        columns = []
                        for col in df.columns:
                            if isinstance(col, pd.Timestamp):
                                columns.append(col.strftime("%Y-%m-%d"))
                            else:
                                columns.append(str(col))

                        df.columns = columns

                        # Write data
                        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
                            for c_idx, value in enumerate(row, 1):
                                if pd.isna(value):
                                    value = ""
                                ws.cell(row=r_idx, column=c_idx, value=value)

                write_df_to_sheet(wb, "7.TIKR_IS", is_df)
                write_df_to_sheet(wb, "8.TIKR_BS", bs_df)
                write_df_to_sheet(wb, "9.TIKR_CF", cf_df)

                output = io.BytesIO()
                wb.save(output)
                output.seek(0)

            st.success("¡Excel generado exitosamente!")

            st.download_button(
                label="📥 Descargar Plantilla Completada",
                data=output,
                file_name=f"Valoracion_IDC_{ticker_symbol}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"Se produjo un error al procesar los datos: {e}")
            logging.error(f"Error processing {ticker_symbol}: {e}", exc_info=True)
    else:
        st.warning("Por favor, introduce un Ticker válido.")
